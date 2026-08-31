#!/usr/bin/env python3
"""Generate ATS-friendly resume targeted at remote hands / smart hands roles."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Amlen_Singha_Remote_Hands_Resume.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=4, line=1.0)
    run = p.add_run(text.upper())
    set_run_font(run, size=12, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    add_bottom_border(p)
    return p


def body(doc, text, size=11, italic=False, bold=False, after=4):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=after, line=1.08)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    return p


def job_header(doc, title, dates):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=0, line=1.05)
    r1 = p.add_run(title)
    set_run_font(r1, size=11, bold=True)
    r2 = p.add_run("\t" + dates)
    set_run_font(r2, size=10.5, bold=False, color=RGBColor(0x33, 0x33, 0x33))
    # tab stop right
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(7.1), alignment=2)  # right-ish
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=2, line=1.05)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=2)
    r = name.add_run("AMLEN SINGHA")
    set_run_font(r, size=20, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    target = doc.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(target, before=0, after=2)
    r = target.add_run(
        "Remote Hands / Smart Hands Technician  |  Data Center Customer Operations"
    )
    set_run_font(r, size=11, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=2)
    r = contact.add_run(
        "India  |  +91-9911989661  |  amlensingha@gmail.com  |  linkedin.com/in/amlensingha"
    )
    set_run_font(r, size=10)

    avail = doc.add_paragraph()
    avail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(avail, before=0, after=6)
    r = avail.add_run(
        "Availability: Immediate / 15 days  ·  On-site  ·  Rotational 24x7 shifts  ·  Physical DC floor work"
    )
    set_run_font(r, size=10, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    heading(doc, "Professional summary")
    body(
        doc,
        "Operations professional targeting Remote Hands / Smart Hands and Data Center Customer "
        "Operations roles. 8+ years executing ticket-style work against SOPs, with documented "
        "accuracy (98%+ QA), vendor coordination, incident follow-through, and time-bound "
        "reporting. Comfortable following customer runbooks, logging work in tickets, and "
        "closing requests to SLA. Linux fundamentals in progress (PGDCA); Japanese bilingual "
        "(JLPT N3 certified, N2 trained) for Japan-account colo customers. Open to junior / "
        "trainee technician title to build hands-on rack, cabling, and break-fix skills on site.",
        size=10.5,
        after=4,
    )

    heading(doc, "Core skills (remote hands / colo operations)")
    skills = [
        (
            "Customer operations: ",
            "Ticket-driven requests, SOP / runbook execution, SLA follow-up, quality checks, "
            "escort-style vendor coordination, shift-ready documentation",
        ),
        (
            "Incident & logs: ",
            "Status-log investigation, multi-team remediation, weekly ops reporting, bottleneck tracking",
        ),
        (
            "Asset & vendor tracking: ",
            "Access permissions, renewals, inventory-style records, payment / contract tracking",
        ),
        (
            "Tools: ",
            "Microsoft Excel (advanced reporting), Linux basics, Python (data/reporting), "
            "documentation control, CAPM (structured project tracking)",
        ),
        (
            "Languages: ",
            "English (professional), Japanese (professional / JLPT N3 certified, N2 trained), Hindi (native)",
        ),
    ]
    for label, rest in skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=2, line=1.05)
        r1 = p.add_run("•  " + label)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(rest)
        set_run_font(r2, size=10.5)

    heading(doc, "Professional experience")

    job_header(doc, "Amazon  —  Quality Analyst (Japanese)", "Oct 2024 – Nov 2024")
    body(
        doc,
        "Customer and operations quality role: ticket-style checks, log investigation, vendor/internal coordination.",
        size=10,
        italic=True,
        after=2,
    )
    bullet(
        doc,
        "Executed time-bound quality requests on Japanese-language content; verified phrasing and substance before release (SOP accuracy).",
    )
    bullet(
        doc,
        "Investigated shipment status logs for non-delivery issues; coordinated multi-department remediation and successful re-ships (incident follow-through).",
    )
    bullet(
        doc,
        "Produced weekly operational reports for senior leadership, highlighting process bottlenecks and improvement actions.",
    )
    bullet(
        doc,
        "Worked with global internal teams and external vendors to implement lasting operational fixes—not one-off workarounds.",
    )

    job_header(doc, "Wipro  —  Senior Media Analyst", "Feb 2018 – Oct 2024")
    body(
        doc,
        "High-volume operations and quality role for Japanese corporate clients; documentation, vendors, and SLA-style accuracy.",
        size=10,
        italic=True,
        after=2,
    )
    bullet(
        doc,
        "Completed 100+ Japanese-language quality checks monthly at 98%+ accuracy—same discipline required for smart-hands ticket close-out.",
    )
    bullet(
        doc,
        "Tracked 100+ global vendor subscriptions: access permissions, renewal deadlines, and payment records (asset / vendor control).",
    )
    bullet(
        doc,
        "Maintained project documentation, online access control, and deadline tracking (change and document control).",
    )
    bullet(
        doc,
        "Converted multimedia customer feedback into actionable executive summaries for Japanese stakeholders.",
    )
    bullet(
        doc,
        "Drove continuous-improvement work using IT tools to raise daily throughput; recognized Best Analyst (month/quarter) and Best Team (Q1 2023).",
    )

    job_header(
        doc,
        "Caliber Interconnect Solutions  —  Japanese Translator / Technical Liaison",
        "Jan 2017 – Jul 2017",
    )
    body(
        doc,
        "On-instruction technical work between Japanese clients and Indian engineering teams (PCB).",
        size=10,
        italic=True,
        after=2,
    )
    bullet(
        doc,
        "Captured customer technical requirements and turned them into clear written specs for engineers (runbook / instruction following).",
    )
    bullet(
        doc,
        "Provided consecutive interpretation on live calls and maintained Q&A logs so actions were not lost between offices.",
    )
    bullet(
        doc,
        "Closed communication gaps that delayed delivery—aligned scope, flagged risk, and kept both sides on the same task list.",
    )

    job_header(doc, "IICT Institute  —  Administrator", "Jun 2008 – Dec 2011")
    bullet(
        doc,
        "Ran day-to-day operations: enrollment workflows, records quality control, scheduling, and stakeholder coordination.",
    )

    heading(doc, "Career break & current availability")
    job_header(
        doc,
        "Japanese language study — Shintomi International Language Institute, Chiba, Japan",
        "Jan 2025 – Mar 2026",
    )
    bullet(
        doc,
        "Full-time Japanese study in Japan (culture, spoken/written Japanese, Japan-style work communication). Now fully available for on-site DC shifts in India.",
    )
    job_header(
        doc,
        "Visual Japanese — Independent learning-content project (part-time, not a full-time employer)",
        "Apr 2026 – Present",
    )
    bullet(
        doc,
        "Side project (YouTube / study materials). Job search and on-site technician roles are the priority; this is not a competing full-time commitment.",
    )

    heading(doc, "Education")
    bullet(
        doc,
        "Post Graduate Diploma in Computer Applications (PGDCA) — IGNOU (pursuing) — building IT and systems fundamentals for data center floor work.",
    )
    bullet(doc, "Bachelor of Arts (B.A.) — Indira Gandhi National Open University (IGNOU), 2008")
    bullet(
        doc,
        "Japanese language: Shintomi (Japan) 2025–2026; MIKADO 2015–2016; MOSAI 2013–2014",
    )

    heading(doc, "Certifications")
    bullet(doc, "CAPM — Certified Associate in Project Management (PMI) — structured tracking, documentation, change control")
    bullet(doc, "JLPT N3 — Japanese Language Proficiency Test (certified); JLPT N2 trained")
    bullet(doc, "IBM: Python for Data Science, AI & Development; Introduction to Data Analytics")
    bullet(doc, "Google AI Professional Certificate")

    heading(doc, "Honors")
    body(
        doc,
        "Wipro: Best Analyst of the Month (May 2024, Oct 2023)  ·  Best Analyst of the Quarter (Q2 2023)  ·  Best Team Award (Q1 2023)",
        size=10.5,
        after=4,
    )

    heading(doc, "What I am ready to do on day one")
    body(
        doc,
        "On-site remote hands / smart hands tickets: follow customer instructions, photograph/confirm work, "
        "update the ticket, meet the SLA, and escalate when the runbook stops. Physical work, nights, and "
        "weekends accepted. Hands-on racking, copper/fiber, and break-fix to be built on the job (trainee / "
        "junior technician). Japanese available for Japan-account customers.",
        size=10.5,
        after=0,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
